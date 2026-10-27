"""Builds small synthetic workbooks with the same layout as the real Chosen One workbook."""

# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false
# pyright: reportUnknownArgumentType=false, reportMissingTypeStubs=false

from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

MONTHS = [date(2026, 10, 1), date(2026, 11, 1), date(2026, 12, 1)]
HEADERS = ["Customer", "Project", "Client Partner", "AE", "TECH", "Region", "Industry", "CATG"]
EXTRA = [
    "Row ID",
    "Commercial Owner",
    "Country",
    "Billing Type",
    "Status",
    "Last Updated",
    "Last Updated By",
]
DEFAULT_USERS = [
    ("AC", "Avery Chen", "Client Partner", "avery.chen@example.com", "Canada", None),
    ("JM", "Jordan Malik", "Client Partner", "jordan.malik@example.com", "US", None),
    ("Shah", "Riley Shah", "Seller / AE", "riley.shah@example.com", "Canada", None),
    ("AF", "Alex Fraser", "Commercial Owner", "alex.fraser@example.com", "Canada", None),
    ("PN", "Priya Nair", "Client Partner; Commercial Owner", "priya@example.com", "India", None),
]


def make_row(row_id: str, **overrides: Any) -> dict[str, Any]:
    row: dict[str, Any] = {
        "Customer": f"Customer {row_id}",
        "Project": f"Project {row_id}",
        "Client Partner": "AC",
        "AE": "N/A",
        "TECH": "APPIAN",
        "Region": "NA",
        "Industry": "Insurance",
        "CATG": "EE",
        "Row ID": row_id,
        "Commercial Owner": None,
        "Country": "Canada",
        "Billing Type": "T&M",
        "Status": "Active",
        "Last Updated": datetime(2026, 9, 14),
        "Last Updated By": "AC",
        "amounts": [1000.0, 2000.0, 3000.0],
    }
    row.update(overrides)
    return row


class WorkbookBuilder:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.rows: list[dict[str, Any]] = []
        self.months = list(MONTHS)
        self.users: list[tuple[Any, ...]] = list(DEFAULT_USERS)

    def add(self, row_id: str, **overrides: Any) -> WorkbookBuilder:
        self.rows.append(make_row(row_id, **overrides))
        return self

    def build(self) -> Path:
        wb = openpyxl.Workbook()
        readme = wb.active
        assert readme is not None
        readme.title = "README"
        readme.append(["The Chosen One - synthetic test data"])
        readme.append(["As of", "2026-10-01. Test workbook."])

        sheet = wb.create_sheet("Forecast")
        first_month_col = len(HEADERS) + 1
        subtotal_row: list[Any] = [None] * (first_month_col - 2) + ["Subtotal (filtered) ->"]
        last_row = len(self.rows) + 2
        for index in range(len(self.months)):
            letter = get_column_letter(first_month_col + index)
            subtotal_row.append(f"=SUBTOTAL(9,{letter}3:{letter}{last_row})")
        sheet.append(subtotal_row)
        sheet.append(HEADERS + [datetime(m.year, m.month, m.day) for m in self.months] + EXTRA)
        for row in self.rows:
            amounts = row["amounts"]
            values = [row[h] for h in HEADERS] + list(amounts)
            values += [None] * (len(self.months) - len(amounts))
            sheet.append(values + [row[h] for h in EXTRA])

        users = wb.create_sheet("Users")
        users.append(["Code", "Name", "Role(s)", "Email", "Location", "Notes"])
        for user in self.users:
            users.append(list(user))
        wb.save(self.path)
        return self.path
