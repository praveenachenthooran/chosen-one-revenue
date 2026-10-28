"""Reporting API: permissions, snapshots, exports, actuals uploads and the accuracy report."""

import io
from datetime import date
from typing import Any

import openpyxl
import pytest
from httpx import AsyncClient, Response
from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import ActualsUpload, ExportEvent, ForecastSnapshot, MonthlyActual, Project
from tests.conftest import SignedIn
from tests.pipeline_helpers import sdeal
from tests.reporting_helpers import MON_NOV_2, csv_file, project_id, set_amount, take, xlsx_file
from tests.test_pipeline_api import Env, env, get  # noqa: F401  # pyright: ignore[reportUnusedImport]

OCT, NOV = date(2026, 10, 1), date(2026, 11, 1)
FINANCE, ADMIN, CRO = "fin@corp.test", "admin@corp.test", "cro@corp.test"
OTHERS = ["avery@corp.test", "riley@corp.test", "res@corp.test", "sam@corp.test"]
CSV = {"Content-Type": "text/csv"}
FILTER_CASES: list[dict[str, str]] = [
    {"owner": "JM"},
    {"practice": "CRM"},
    {"region": "UK"},
    {"industry": "Insurance"},
    {"category": "EN"},
    {"customer": "initech"},
    {"owner": "Shah", "region": "UK"},
    {"q": "mule"},
]


async def post_snapshot(
    http: AsyncClient, headers: dict[str, str], note: str | None = None
) -> Response:
    return await http.post("/v1/snapshots", json={"note": note} if note else {}, headers=headers)


async def send_actuals(
    http: AsyncClient,
    headers: dict[str, str],
    content: bytes,
    filename: str = "actuals.csv",
    **params: Any,
) -> Response:
    return await http.post(
        "/v1/actuals/uploads",
        params={"filename": filename, **params},
        content=content,
        headers={**headers, "Content-Type": "application/octet-stream"},
    )


def sheet_rows(content: bytes, name: str) -> list[list[Any]]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    return [list(r) for r in workbook[name].iter_rows(values_only=True)]


def summary(content: bytes) -> dict[str, str]:
    return {str(r[0]): str(r[1]) for r in sheet_rows(content, "Summary") if r[0] and r[1]}


# ---- access -------------------------------------------------------------------------------

READS = [
    "/v1/snapshots",
    "/v1/snapshots/1",
    "/v1/snapshots/1/verify",
    "/v1/snapshots/1/export",
    "/v1/forecast/export",
    "/v1/actuals/uploads",
    "/v1/actuals/uploads/1",
    "/v1/actuals/months",
    "/v1/accuracy",
]


async def test_every_reporting_endpoint_requires_sign_in(client: AsyncClient) -> None:
    for path in READS:
        assert (await client.get(path)).status_code == 401, path
    assert (await client.post("/v1/snapshots", json={})).status_code == 401
    assert (
        await client.post("/v1/actuals/uploads", params={"filename": "a.csv"})
    ).status_code == 401


async def test_mutations_need_the_csrf_token(env: Env, as_user: SignedIn) -> None:
    http, _ = await as_user(ADMIN)
    responses = [
        await http.post("/v1/snapshots", json={}),
        await http.post("/v1/actuals/uploads", params={"filename": "a.csv"}, content=b"x"),
    ]
    for response in responses:
        assert response.status_code == 403, response.url
        assert response.json()["detail"]["code"] == "csrf_failed"
    assert await env.session.scalar(select(func.count()).select_from(ForecastSnapshot)) == 0


@pytest.mark.parametrize("email", [FINANCE, ADMIN])
async def test_finance_and_admin_can_take_a_snapshot(
    env: Env, as_user: SignedIn, email: str
) -> None:
    http, headers = await as_user(email)
    response = await post_snapshot(http, headers, "Board pack")
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["kind"] == "manual" and body["note"] == "Board pack"
    assert body["combined_total"] == "18000.00" and len(body["checksum"]) == 64
    assert body["taken_by_label"] in ("Fin Ance", "Ada Admin")


@pytest.mark.parametrize("email", [CRO, *OTHERS])
async def test_nobody_else_can_take_a_snapshot(env: Env, as_user: SignedIn, email: str) -> None:
    http, headers = await as_user(email)
    response = await post_snapshot(http, headers)
    assert response.status_code == 403 and response.json()["detail"]["code"] == "forbidden_role"
    assert await env.session.scalar(select(func.count()).select_from(ForecastSnapshot)) == 0


@pytest.mark.parametrize("email", [FINANCE, ADMIN, CRO, *OTHERS])
async def test_every_signed_in_role_can_read_and_export_snapshots(
    env: Env, as_user: SignedIn, email: str
) -> None:
    await take(env.session, MON_NOV_2)
    http, _ = await as_user(email)
    assert (await get(http, "/v1/snapshots"))["can_take"] is (email in (FINANCE, ADMIN))
    assert (await get(http, "/v1/snapshots/1"))["snapshot"]["id"] == 1
    assert (await get(http, "/v1/snapshots/1/verify"))["intact"] is True
    assert (await http.get("/v1/snapshots/1/export")).status_code == 200
    assert (await http.get("/v1/forecast/export")).status_code == 200


@pytest.mark.parametrize("email", [FINANCE, ADMIN, CRO])
async def test_finance_leadership_and_admin_can_see_actuals_and_accuracy(
    env: Env, as_user: SignedIn, email: str
) -> None:
    http, _ = await as_user(email)
    for path in ("/v1/actuals/uploads", "/v1/actuals/months", "/v1/accuracy"):
        assert (await http.get(path)).status_code == 200, path


@pytest.mark.parametrize("email", OTHERS)
async def test_other_roles_cannot_see_actuals_or_accuracy(
    env: Env, as_user: SignedIn, email: str
) -> None:
    http, headers = await as_user(email)
    for path in ("/v1/actuals/uploads", "/v1/actuals/months", "/v1/accuracy"):
        response = await http.get(path)
        assert response.status_code == 403 and response.json()["detail"]["code"] == "forbidden_role"
    assert (
        await send_actuals(http, headers, csv_file("2026-10,Kestrel Life,,1"))
    ).status_code == 403


async def test_only_finance_and_admin_can_load_actuals(env: Env, as_user: SignedIn) -> None:
    content = csv_file("2026-10,Kestrel Life,Policy Admin,1000")
    cro, cro_headers = await as_user(CRO)
    assert (await send_actuals(cro, cro_headers, content)).status_code == 403
    assert await env.session.scalar(select(func.count()).select_from(MonthlyActual)) == 0
    for email in (FINANCE, ADMIN):
        http, headers = await as_user(email)
        response = await send_actuals(http, headers, content, dry_run="true")
        assert response.status_code == 200, response.text


async def test_unknown_snapshot_and_upload_are_404(env: Env, as_user: SignedIn) -> None:
    http, _ = await as_user(FINANCE)
    for path in (
        "/v1/snapshots/999",
        "/v1/snapshots/999/verify",
        "/v1/snapshots/999/export",
        "/v1/actuals/uploads/999",
    ):
        response = await http.get(path)
        assert response.status_code == 404, path


# ---- snapshots through the API ------------------------------------------------------------


async def test_snapshot_totals_match_the_live_forecast_at_that_moment(
    env: Env, as_user: SignedIn
) -> None:
    await env.load([sdeal("1", company="Initech", practice="CRM", region="UK")])
    http, headers = await as_user(FINANCE)
    live = await get(http, "/v1/forecast/monthly")
    created = (await post_snapshot(http, headers)).json()
    stored = await get(http, f"/v1/snapshots/{created['id']}")

    assert stored["months"] == live["months"]
    assert stored["fiscal_years"] == live["fiscal_years"]
    for key in ("committed_total", "weighted_total", "combined_total"):
        assert stored[key] == live[key]
    assert stored["unphased"] == live["unphased"]
    assert stored["outside_window_weighted"] == live["outside_window_weighted"]
    assert stored["filters_applied"] is False and stored["snapshot"]["id"] == created["id"]


@pytest.mark.parametrize("filters", FILTER_CASES, ids=[str(f) for f in FILTER_CASES])
async def test_filtered_snapshot_equals_the_filtered_live_forecast(
    env: Env, as_user: SignedIn, filters: dict[str, str]
) -> None:
    await env.load(
        [
            sdeal("1", company="Initech", practice="CRM", region="UK"),
            sdeal("2", company="Granitepeak Telecom", owner="Jordan Malik", amount="50000"),
        ]
    )
    http, headers = await as_user(FINANCE)
    live = await get(http, "/v1/forecast/monthly", **filters)
    snapshot_id = (await post_snapshot(http, headers)).json()["id"]
    stored = await get(http, f"/v1/snapshots/{snapshot_id}", **filters)

    assert stored["months"] == live["months"], filters
    assert stored["combined_total"] == live["combined_total"] and stored["filters_applied"] is True
    assert stored["unphased"] is None and stored["outside_window_weighted"] is None


async def test_a_snapshot_does_not_move_when_the_live_forecast_is_edited(
    env: Env, as_user: SignedIn
) -> None:
    http, headers = await as_user(ADMIN)
    created = (await post_snapshot(http, headers)).json()
    before = await get(http, f"/v1/snapshots/{created['id']}")

    pid = await project_id(env.session, "R1")
    edit = await http.put(
        f"/v1/projects/{pid}/revenue/2026-10-01",
        json={"amount": "9999", "expected_amount": "1000.00"},
        headers=headers,
    )
    assert edit.status_code == 200, edit.text
    live = await get(http, "/v1/forecast/monthly")
    assert live["combined_total"] != before["combined_total"]

    assert await get(http, f"/v1/snapshots/{created['id']}") == before
    assert (await get(http, f"/v1/snapshots/{created['id']}/verify"))["intact"] is True
    for method, path in (
        ("put", "/v1/snapshots/1"),
        ("delete", "/v1/snapshots/1"),
        ("patch", "/v1/snapshots/1"),
    ):
        assert (
            await getattr(http, method)(path, headers=headers)
        ).status_code == 405  # no way to change one


async def test_the_snapshot_list_shows_history_and_the_schedule(
    env: Env, as_user: SignedIn
) -> None:
    http, headers = await as_user(FINANCE)
    empty = await get(http, "/v1/snapshots")
    assert empty["items"] == [] and empty["scheduled_taken_this_month"] is False
    assert empty["last_job"] is None and date.fromisoformat(empty["next_scheduled"]).weekday() < 5
    await post_snapshot(http, headers)
    await post_snapshot(http, headers, "second")
    body = await get(http, "/v1/snapshots")
    assert [s["note"] for s in body["items"]] == ["second", None]  # newest first


async def test_edits_to_snapshotted_months_are_flagged_in_the_history(
    env: Env, as_user: SignedIn
) -> None:
    http, headers = await as_user(ADMIN)
    pid = await project_id(env.session, "R1")

    async def edit(month: str, amount: str, expected: str) -> None:
        response = await http.put(
            f"/v1/projects/{pid}/revenue/{month}",
            json={"amount": amount, "expected_amount": expected},
            headers=headers,
        )
        assert response.status_code == 200, response.text

    await edit("2026-10-01", "1100", "1000.00")  # before any snapshot
    await post_snapshot(http, headers)
    await edit(
        "2026-11-01", "2100", "2000.00"
    )  # after: the snapshot is unchanged, the edit flagged

    items = (await get(http, f"/v1/projects/{pid}/history"))["items"]
    flags = {(i["month"], i["new_value"]): (i["post_snapshot"], i["from_import"]) for i in items}
    assert flags[("2026-10-01", "1100.00")] == (False, False)
    assert flags[("2026-11-01", "2100.00")] == (True, False)
    assert all(not i["post_snapshot"] for i in items if i["from_import"])


# ---- exports ------------------------------------------------------------------------------


async def test_current_forecast_xlsx_matches_the_on_screen_totals(
    env: Env, as_user: SignedIn
) -> None:
    await env.load([sdeal("1")])
    http, _ = await as_user(FINANCE)
    live = await get(http, "/v1/forecast/monthly")
    response = await http.get("/v1/forecast/export", params={"format": "xlsx"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    info = summary(response.content)
    assert info["Export"] == "Current forecast" and info["Exported by"] == "Fin Ance"
    assert info["Filters applied"] == "None (all data)" and "Exported at (UTC)" in info
    assert info["Currency"].startswith("Amounts as stored")

    rows = sheet_rows(response.content, "Monthly")
    assert rows[0] == [
        "Month",
        "Fiscal year",
        "Fiscal period",
        "Committed",
        "Weighted",
        "Total forecast",
    ]
    by_month = {r[0]: r for r in rows[1:] if r[1] and r[2]}
    for month in live["months"]:
        row = by_month[month["month"][:7]]
        assert (f"{row[3]:.2f}", f"{row[4]:.2f}", f"{row[5]:.2f}") == (
            month["committed"],
            month["weighted"],
            month["total"],
        )
    grand = rows[-1]
    assert grand[0] == "Grand total" and f"{grand[5]:.2f}" == live["combined_total"]
    assert (
        f"{grand[3]:.2f}" == live["committed_total"] and f"{grand[4]:.2f}" == live["weighted_total"]
    )
    assert {"Summary", "Monthly", "Projects", "Pipeline"} == set(
        openpyxl.load_workbook(io.BytesIO(response.content)).sheetnames
    )


@pytest.mark.parametrize("filters", FILTER_CASES, ids=[str(f) for f in FILTER_CASES])
async def test_filtered_export_states_its_filters_and_matches_the_filtered_screen(
    env: Env, as_user: SignedIn, filters: dict[str, str]
) -> None:
    await env.load([sdeal("1", company="Initech", practice="CRM", region="UK")])
    http, _ = await as_user(FINANCE)
    live = await get(http, "/v1/forecast/monthly", **filters)
    response = await http.get("/v1/forecast/export", params={"format": "xlsx", **filters})
    info = summary(response.content)

    assert info["Filters applied"] != "None (all data)"
    for value in filters.values():
        assert value in info["Filters applied"]
    grand = sheet_rows(response.content, "Monthly")[-1]
    assert f"{grand[5]:.2f}" == live["combined_total"], filters


async def test_snapshot_export_carries_the_stored_figures_and_snapshot_metadata(
    env: Env, as_user: SignedIn
) -> None:
    await env.load([sdeal("1")])
    http, headers = await as_user(FINANCE)
    created = (await post_snapshot(http, headers, "Board")).json()
    await set_amount(env.session, "R1", OCT, "50000.00")  # live data moves on

    response = await http.get(f"/v1/snapshots/{created['id']}/export", params={"format": "xlsx"})
    info = summary(response.content)
    assert info["Export"] == f"Forecast snapshot #{created['id']}"
    assert info["Snapshot checksum (SHA-256)"] == created["checksum"]
    assert info["Snapshot kind"] == "manual" and "Snapshot taken (UTC)" in info
    assert info["Snapshot taken by"] == "Fin Ance" and "Exported at (UTC)" in info
    grand = sheet_rows(response.content, "Monthly")[-1]
    assert f"{grand[5]:.2f}" == created["combined_total"] == "78000.00"  # not the edited figure


async def test_csv_export_has_metadata_a_bom_and_each_view(env: Env, as_user: SignedIn) -> None:
    await env.load([sdeal("1")])
    http, _ = await as_user(FINANCE)
    raw = (await http.get("/v1/forecast/export", params={"format": "csv", "owner": "JM"})).content
    assert raw.startswith(b"\xef\xbb\xbf")
    lines = raw.decode("utf-8-sig").splitlines()
    assert lines[0] == "Export,Current forecast"
    assert any(line.startswith("Exported at (UTC),") for line in lines)
    assert any(line.startswith("Filters applied,Owner (user code) = JM") for line in lines)
    assert "View,Monthly" in lines and lines[lines.index("View,Monthly") + 1] == ""
    header_at = lines.index("View,Monthly") + 2
    assert lines[header_at] == "Month,Fiscal year,Fiscal period,Committed,Weighted,Total forecast"
    assert lines[header_at + 1].startswith("2026-10,FY27,7,")

    projects = (
        await http.get("/v1/forecast/export", params={"format": "csv", "view": "projects"})
    ).text
    assert "Row ID,Customer,Project,Status,TECH,Region,Industry,Category,Owners" in projects
    assert "R1,Kestrel Life,Policy Admin" in projects
    pipeline = (
        await http.get("/v1/forecast/export", params={"format": "csv", "view": "pipeline"})
    ).text
    assert (
        "HubSpot Deal ID,Deal,Company,Stage,Probability" in pipeline
        and "Negotiation/Review" in pipeline
    )

    response = await http.get("/v1/forecast/export", params={"format": "csv", "view": "bogus"})
    assert response.status_code == 422 and response.json()["detail"]["code"] == "invalid_view"
    assert (await http.get("/v1/forecast/export", params={"format": "pdf"})).status_code == 422


async def test_a_name_that_looks_like_a_formula_is_exported_as_text(
    env: Env, as_user: SignedIn
) -> None:
    await env.session.execute(
        update(Project)
        .where(Project.source_row_id == "R1")
        .values(name='=HYPERLINK("http://evil","x")')
    )
    await env.session.commit()
    http, _ = await as_user(FINANCE)
    xlsx = (await http.get("/v1/forecast/export", params={"format": "xlsx"})).content
    names = [r[2] for r in sheet_rows(xlsx, "Projects")[1:]]
    assert '\'=HYPERLINK("http://evil","x")' in names
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert all(cell.data_type != "f" for row in workbook["Projects"].iter_rows() for cell in row)
    csv_text = (
        await http.get("/v1/forecast/export", params={"format": "csv", "view": "projects"})
    ).text
    assert "'=HYPERLINK(" in csv_text and ",=HYPERLINK" not in csv_text


async def test_exports_are_logged_with_who_and_which_filters(env: Env, as_user: SignedIn) -> None:
    http, headers = await as_user(FINANCE)
    snapshot_id = (await post_snapshot(http, headers)).json()["id"]
    await http.get("/v1/forecast/export", params={"format": "csv", "region": "UK"})
    await http.get(f"/v1/snapshots/{snapshot_id}/export", params={"format": "xlsx"})
    events = list(await env.session.scalars(select(ExportEvent).order_by(ExportEvent.id)))
    assert [(e.source, e.file_format, e.snapshot_id, e.filters) for e in events] == [
        ("current", "csv", None, {"region": "UK"}),
        ("snapshot", "xlsx", snapshot_id, {}),
    ]
    assert all(e.exported_by_id is not None for e in events)


# ---- actuals through the API --------------------------------------------------------------


async def test_actuals_upload_returns_the_reconciliation_and_is_recorded(
    env: Env, as_user: SignedIn
) -> None:
    http, headers = await as_user(FINANCE)
    content = csv_file(
        "2026-10,Kestrel Life,Policy Admin,1000", "2026-10,Initech,Existing Work,1500"
    )
    response = await send_actuals(http, headers, content, "october-actuals.csv")

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "completed" and body["dry_run"] is False and body["upload_id"]
    report = body["report"]
    assert report["total_amount"] == "2500.00" and report["rows_new"] == 2
    month = report["months"][0]
    assert (month["actual_total"], month["committed_forecast_now"]) == ("2500.00", "3000.00")

    uploads = await get(http, "/v1/actuals/uploads")
    assert [(u["filename"], u["status"], u["uploaded_by_label"]) for u in uploads] == [
        ("october-actuals.csv", "completed", "Fin Ance")
    ]
    detail = await get(http, f"/v1/actuals/uploads/{body['upload_id']}")
    assert detail["report"]["total_amount"] == "2500.00"
    months = await get(http, "/v1/actuals/months")
    assert months == [
        {
            "month": "2026-10-01",
            "fiscal_year": "FY27",
            "actual_total": "2500.00",
            "committed_forecast": "3000.00",
        }
    ]


async def test_a_bad_file_is_rejected_with_every_problem_and_stores_nothing(
    env: Env, as_user: SignedIn
) -> None:
    http, headers = await as_user(FINANCE)
    bad = csv_file(
        "2026-10,Kestrel Life,Policy Admin,1000",
        "2026-10,Nobody,,1",
        "2026-10,Kestrel Life,Policy Admin,abc",
    )
    response = await send_actuals(http, headers, bad)
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail["code"] == "actuals_rejected" and "2 row(s)" in detail["message"]
    assert {e["code"] for e in detail["report"]["errors"]} == {"unknown_customer", "invalid_amount"}
    assert [e["row"] for e in detail["report"]["errors"]] == [3, 4]
    assert await env.session.scalar(select(func.count()).select_from(MonthlyActual)) == 0
    attempt = await env.session.scalar(select(ActualsUpload))
    assert (
        attempt is not None and attempt.status == "rejected" and attempt.id == detail["upload_id"]
    )


@pytest.mark.parametrize(
    ("filename", "content", "status", "code"),
    [
        ("a.csv", b"", 400, "empty_file"),
        ("a.xls", b"x", 415, "unsupported_file_type"),
        ("a.xlsx", b"nope", 400, "unreadable_file"),
        ("a.csv", b"Month,Amount\n2026-10,1\n", 400, "missing_columns"),
    ],
)
async def test_unusable_files_get_a_clear_error(
    env: Env, as_user: SignedIn, filename: str, content: bytes, status: int, code: str
) -> None:
    http, headers = await as_user(FINANCE)
    response = await send_actuals(http, headers, content, filename)
    assert response.status_code == status, response.text
    assert response.json()["detail"]["code"] == code


async def test_oversized_upload_is_refused(env: Env, as_user: SignedIn) -> None:
    http, headers = await as_user(FINANCE)
    response = await send_actuals(http, headers, b"x" * (5 * 1024 * 1024 + 10))
    assert response.status_code == 413 and response.json()["detail"]["code"] == "file_too_large"


async def test_a_missing_filename_is_a_validation_error(env: Env, as_user: SignedIn) -> None:
    http, headers = await as_user(FINANCE)
    response = await http.post("/v1/actuals/uploads", content=b"x", headers=headers)
    assert response.status_code == 422


async def test_xlsx_upload_dry_run_and_repeat_through_the_api(env: Env, as_user: SignedIn) -> None:
    http, headers = await as_user(ADMIN)
    content = xlsx_file([["2026-10", "Kestrel Life", "Policy Admin", 1000]])
    preview = await send_actuals(http, headers, content, "a.xlsx", dry_run="true")
    assert preview.json()["dry_run"] is True and preview.json()["upload_id"] is None
    assert await env.session.scalar(select(func.count()).select_from(MonthlyActual)) == 0

    first = await send_actuals(http, headers, content, "a.xlsx")
    again = await send_actuals(http, headers, content, "a.xlsx")
    assert first.json()["report"]["rows_new"] == 1
    assert (again.json()["report"]["rows_new"], again.json()["report"]["rows_unchanged"]) == (0, 1)
    assert await env.session.scalar(select(func.count()).select_from(MonthlyActual)) == 1


# ---- the accuracy report through the API --------------------------------------------------


async def test_accuracy_report_end_to_end(env: Env, as_user: SignedIn) -> None:
    from datetime import UTC, datetime

    from app.snapshots.service import SCHEDULED

    await take(env.session, datetime(2026, 7, 1, 9, tzinfo=UTC), kind=SCHEDULED)
    await set_amount(env.session, "R1", OCT, "1200.00")
    await take(env.session, datetime(2026, 9, 1, 9, tzinfo=UTC), kind=SCHEDULED)

    http, headers = await as_user(FINANCE)
    upload = await send_actuals(http, headers, csv_file("2026-10,Kestrel Life,,3000"))
    assert upload.status_code == 200, upload.text
    report = await get(http, "/v1/accuracy")

    assert report["tolerance_percent"] == "5" and report["target_lag"] == 3
    (october,) = report["months"]
    lags = {entry["lag"]: entry for entry in october["lags"]}
    assert (lags[3]["forecast"], lags[3]["variance_percent"], lags[3]["flagged"]) == (
        "3000.00",
        "0.00",
        False,
    )
    assert lags[2]["status"] == "no_snapshot"
    assert (lags[1]["forecast"], lags[1]["variance_percent"], lags[1]["flagged"]) == (
        "3200.00",
        "6.67",
        True,
    )
    totals = {t["lag"]: t for t in report["totals"]}
    assert totals[1]["months_flagged"] == 1 and totals[1]["flagged"] is True
    assert totals[3]["months_compared"] == 1 and totals[3]["variance_percent"] == "0.00"


async def test_accuracy_is_empty_before_any_actuals(env: Env, as_user: SignedIn) -> None:
    http, _ = await as_user(FINANCE)
    report = await get(http, "/v1/accuracy")
    assert report["months"] == [] and [t["months_compared"] for t in report["totals"]] == [0, 0, 0]


async def test_reporting_failures_do_not_affect_the_existing_forecast(
    env: Env, as_user: SignedIn, session: AsyncSession
) -> None:
    http, headers = await as_user(FINANCE)
    before = await get(http, "/v1/forecast/monthly")
    await send_actuals(http, headers, b"garbage", "a.xlsx")
    await send_actuals(http, headers, csv_file("2026-10,Nobody,,1"))
    assert (await http.get("/v1/snapshots/999")).status_code == 404
    assert await get(http, "/v1/forecast/monthly") == before
    assert (await http.get("/v1/projects")).status_code == 200
