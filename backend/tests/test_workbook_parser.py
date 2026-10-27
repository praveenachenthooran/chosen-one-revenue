"""Parser and validation tests. No database needed."""

# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false
# pyright: reportUnknownArgumentType=false, reportMissingTypeStubs=false

import hashlib
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.migration.workbook import ParsedWorkbook, WorkbookStructureError, parse_workbook
from tests.workbook_builder import MONTHS, WorkbookBuilder


def codes(parsed: ParsedWorkbook, *, rejected: bool | None = None) -> list[tuple[str, str | None]]:
    return [
        (i.code, i.source_row_id)
        for i in parsed.issues
        if rejected is None or i.row_rejected == rejected
    ]


def test_valid_rows_are_parsed_with_exact_decimals(builder: WorkbookBuilder) -> None:
    builder.add("R1", amounts=[1234567.89, None, 0.1])
    parsed = parse_workbook(builder.build())

    assert parsed.rows_read == 1
    assert parsed.rejected_row_count == 0
    project = parsed.projects[0]
    assert project.source_row_id == "R1"
    assert project.months == {MONTHS[0]: Decimal("1234567.89"), MONTHS[2]: Decimal("0.10")}
    assert project.cells[MONTHS[0]] == "Forecast!I3"
    assert parsed.month_columns == MONTHS
    assert parsed.as_of == date(2026, 10, 1)


def test_blank_amount_is_unknown_not_zero(builder: WorkbookBuilder) -> None:
    builder.add("R1", amounts=[100.0, None, 300.0])
    project = parse_workbook(builder.build()).projects[0]
    assert MONTHS[1] not in project.months


def test_fully_blank_rows_are_skipped_not_rejected(builder: WorkbookBuilder) -> None:
    builder.add("R1").add("R2")
    path = builder.build()
    wb = openpyxl.load_workbook(path)
    for cell in wb["Forecast"][3]:  # blank out the R1 row, leaving a gap before R2
        cell.value = None
    wb.save(path)
    parsed = parse_workbook(path)
    assert [p.source_row_id for p in parsed.projects] == ["R2"]
    assert parsed.blank_rows == 1 and parsed.rejected_row_count == 0


@pytest.mark.parametrize(
    ("bad_value", "code"),
    [
        ("N/A", "invalid_amount"),
        ("1,200", "invalid_amount"),
        (-50.0, "negative_amount"),
        (10.005, "amount_precision"),
        (True, "invalid_amount"),
    ],
)
def test_invalid_amounts_reject_the_row(
    builder: WorkbookBuilder, bad_value: object, code: str
) -> None:
    builder.add("R1").add("R2", amounts=[1.0, bad_value, 3.0])
    parsed = parse_workbook(builder.build())

    assert [p.source_row_id for p in parsed.projects] == ["R1"]
    assert (code, "R2") in codes(parsed, rejected=True)
    assert parsed.rejected_row_count == 1


def test_formula_cells_are_rejected_never_evaluated(builder: WorkbookBuilder) -> None:
    builder.add("R1", amounts=["=90000-55000", 2.0, 3.0])
    parsed = parse_workbook(builder.build())
    assert parsed.projects == []
    assert ("formula_in_amount", "R1") in codes(parsed, rejected=True)


@pytest.mark.parametrize("field", ["Customer", "Project", "Row ID", "Status"])
def test_missing_required_fields_reject_the_row(builder: WorkbookBuilder, field: str) -> None:
    builder.add("R1", **{field: None})
    parsed = parse_workbook(builder.build())
    assert parsed.projects == []
    assert any(i.code == "missing_required_field" and i.field == field for i in parsed.issues)


def test_duplicate_row_id_and_duplicate_project_are_rejected(builder: WorkbookBuilder) -> None:
    builder.add("R1", Customer="Acme", Project="Portal")
    builder.add("R1", Customer="Other", Project="Other")  # reused Row ID
    builder.add("R3", Customer=" ACME ", Project="portal")  # same project, different case
    parsed = parse_workbook(builder.build())

    assert [p.source_row_id for p in parsed.projects] == ["R1"]
    found = codes(parsed, rejected=True)
    assert ("duplicate_row_id", "R1") in found
    assert ("duplicate_project", "R3") in found


def test_near_duplicate_project_is_a_warning_and_still_imported(builder: WorkbookBuilder) -> None:
    builder.add("R1", Customer="Lakemoor", Project="Clinical Workflow")
    builder.add("R2", Customer="Lakemoor", Project="Clinical Workflows")
    builder.add("R3", Customer="Lakemoor", Project="Billing Gateway")
    parsed = parse_workbook(builder.build())

    assert len(parsed.projects) == 3
    assert codes(parsed, rejected=False) == [("possible_duplicate_project", "R2")]


def test_invalid_category_is_rejected_and_blank_category_warns(builder: WorkbookBuilder) -> None:
    builder.add("R1", CATG="ZZ").add("R2", CATG=None)
    parsed = parse_workbook(builder.build())
    assert ("invalid_category", "R1") in codes(parsed, rejected=True)
    assert ("missing_category", "R2") in codes(parsed, rejected=False)


def test_na_is_a_valid_region_but_na_slash_owner_means_unassigned(
    builder: WorkbookBuilder,
) -> None:
    builder.add("R1", Region="NA", **{"Client Partner": "N/A"})
    parsed = parse_workbook(builder.build())
    project = parsed.projects[0]
    assert project.region == "NA"  # North America, not "not applicable"
    assert "client_partner" not in project.owner_codes
    assert ("missing_client_partner", "R1") in codes(parsed)


def test_unknown_owner_codes_warn_and_are_not_linked(builder: WorkbookBuilder) -> None:
    builder.add("R1", AE="Ghost", **{"Commercial Owner": "AF"})
    parsed = parse_workbook(builder.build())
    assert parsed.projects[0].owner_codes == {"client_partner": "AC", "commercial_owner": "AF"}
    assert ("unknown_owner_code", "R1") in codes(parsed, rejected=False)


def test_users_are_read_without_trusting_emails(builder: WorkbookBuilder) -> None:
    builder.users.append(("DUP", "Dup Email", "Finance", "avery.chen@example.com", None, None))
    builder.users.append(("BAD", "Bad Email", "Admin", "not-an-email", None, None))
    builder.add("R1")
    parsed = parse_workbook(builder.build())
    by_code = {u.code: u for u in parsed.users}
    assert by_code["PN"].roles == ["Client Partner", "Commercial Owner"]
    assert by_code["DUP"].email is None and by_code["BAD"].email is None
    assert {i.code for i in parsed.issues} >= {"user_email_invalid_or_duplicate"}


def test_source_totals_include_rejected_rows_so_variance_is_visible(
    builder: WorkbookBuilder,
) -> None:
    builder.add("R1", amounts=[100.0, 200.0, 300.0])
    builder.add("R2", Status=None, amounts=[10.0, 20.0, 30.0])
    parsed = parse_workbook(builder.build())
    assert parsed.source_totals[MONTHS[0]] == Decimal("110.00")
    assert sum(p.months.get(MONTHS[0], 0) for p in parsed.projects) == Decimal("100.00")


def test_missing_headers_or_sheet_raise_a_structure_error(
    builder: WorkbookBuilder, tmp_path: Path
) -> None:
    builder.add("R1")
    path = builder.build()
    wb = openpyxl.load_workbook(path)
    wb["Forecast"]["A2"].value = "Client"  # renames the Customer header
    wb.save(path)
    with pytest.raises(WorkbookStructureError):
        parse_workbook(path)

    wb = openpyxl.load_workbook(builder.build())
    del wb["Forecast"]
    wb.save(tmp_path / "no_forecast.xlsx")
    with pytest.raises(WorkbookStructureError, match="Forecast"):
        parse_workbook(tmp_path / "no_forecast.xlsx")


def test_parsing_never_modifies_the_source_file(builder: WorkbookBuilder) -> None:
    builder.add("R1")
    path = builder.build()
    before = (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
    parsed = parse_workbook(path)
    assert (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns) == before
    assert parsed.sha256 == before[0]
