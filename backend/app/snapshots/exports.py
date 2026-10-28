"""Excel and CSV exports of the current forecast or a saved snapshot.

Both come from the same lines, so an export always agrees with the totals on screen. Every export
states when it was made, by whom, what it shows and which filters were applied. Text that a
spreadsheet could read as a formula (starting with = + - @) is prefixed with an apostrophe so a
customer or project name can never run as one. Amounts are exported as stored: there is no
currency in the source data and no conversion is done.
"""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.fiscal import fiscal_period, fiscal_year
from app.pipeline.forecast import DealFilters
from app.snapshots.service import Line, MonthTotal

ZERO = Decimal("0.00")
CSV_VIEWS = ("monthly", "projects", "pipeline")
FILTER_LABELS = (
    ("owner", "Owner (user code)"),
    ("practice", "TECH / practice"),
    ("region", "Region"),
    ("industry", "Industry"),
    ("category", "Category"),
    ("customer", "Customer"),
    ("q", "Search text"),
)
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_TYPE = "text/csv; charset=utf-8"
UNSAFE_START = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value: str | None) -> str:
    """Text for a cell, neutralised if a spreadsheet would treat it as a formula."""
    text = value or ""
    return "'" + text if text.startswith(UNSAFE_START) else text


def filters_dict(filters: DealFilters) -> dict[str, str]:
    return {name: value for name, _ in FILTER_LABELS if (value := getattr(filters, name))}


@dataclass
class ExportData:
    title: str  # e.g. "Current forecast" or "Forecast snapshot #3"
    meta: list[tuple[str, str]]  # source-specific facts (snapshot id, taken at, as of, ...)
    months: list[MonthTotal]
    lines: list[Line]  # already filtered
    filters: DealFilters
    exported_at: datetime
    exported_by: str


@dataclass
class Table:
    name: str
    header: list[str]
    rows: list[list[object]]


def _summary_rows(data: ExportData) -> list[tuple[str, str]]:
    applied = filters_dict(data.filters)
    labels = dict(FILTER_LABELS)
    rows: list[tuple[str, str]] = [
        ("Export", data.title),
        ("Exported at (UTC)", data.exported_at.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S")),
        ("Exported by", data.exported_by),
        *data.meta,
        (
            "Filters applied",
            "; ".join(f"{labels[k]} = {v}" for k, v in applied.items()) or "None (all data)",
        ),
        ("Fiscal year", "April to March (FY27 = Apr 2026 to Mar 2027)"),
        ("Currency", "Amounts as stored; no currency conversion is applied"),
        ("Committed", "Confirmed projects only; unconfirmed Closed Won drafts are excluded"),
        ("Weighted", "Open HubSpot deals at stage probability; never also counted as committed"),
    ]
    return rows


def _monthly(data: ExportData) -> Table:
    header = ["Month", "Fiscal year", "Fiscal period", "Committed", "Weighted", "Total forecast"]
    rows: list[list[object]] = []
    by_year: dict[str, list[MonthTotal]] = defaultdict(list)
    for month in data.months:
        by_year[fiscal_year(month.month)].append(month)
        rows.append(
            [
                month.month.strftime("%Y-%m"), fiscal_year(month.month),
                fiscal_period(month.month), month.committed, month.weighted, month.total,
            ]
        )  # fmt: skip
    for year, group in sorted(by_year.items()):
        rows.append(
            [
                f"{year} total", year, None, sum((m.committed for m in group), ZERO),
                sum((m.weighted for m in group), ZERO), sum((m.total for m in group), ZERO),
            ]
        )  # fmt: skip
    committed = sum((m.committed for m in data.months), ZERO)
    weighted = sum((m.weighted for m in data.months), ZERO)
    rows.append(["Grand total", None, None, committed, weighted, committed + weighted])
    return Table("Monthly", header, rows)


def _wide(name: str, kind: str, data: ExportData, descriptors: list[tuple[str, str]]) -> Table:
    """One row per project (or deal), one column per month."""
    months = [m.month for m in data.months]
    header = [label for _, label in descriptors] + [m.strftime("%Y-%m") for m in months] + ["Total"]
    grouped: dict[str, list[Line]] = defaultdict(list)
    for line in data.lines:
        if line.line_type == kind:
            grouped[line.reference].append(line)
    rows: list[list[object]] = []
    for reference in sorted(grouped):
        lines = grouped[reference]
        first = lines[0]
        amounts: dict[date, Decimal] = defaultdict(lambda: ZERO)
        for line in lines:
            amounts[line.month] += line.amount
        values: dict[str, object] = {
            "reference": safe_text(reference),
            "customer": safe_text(first.customer),
            "name": safe_text(first.name),
            "status": safe_text(first.status),
            "practice": safe_text(first.practice),
            "region": safe_text(first.region),
            "industry": safe_text(first.industry),
            "category": safe_text(first.category),
            "owners": ", ".join(safe_text(c) for c in first.owner_codes),
            "stage": safe_text(first.stage),
            "probability": first.probability,
        }
        row = [values[key] for key, _ in descriptors]
        row += [amounts.get(m, ZERO) for m in months]
        row.append(sum(amounts.values(), ZERO))
        rows.append(row)
    return Table(name, header, rows)


def tables(data: ExportData) -> dict[str, Table]:
    return {
        "monthly": _monthly(data),
        "projects": _wide(
            "Projects", "committed", data,
            [
                ("reference", "Row ID"), ("customer", "Customer"), ("name", "Project"),
                ("status", "Status"), ("practice", "TECH"), ("region", "Region"),
                ("industry", "Industry"), ("category", "Category"), ("owners", "Owners"),
            ],
        ),
        "pipeline": _wide(
            "Pipeline", "weighted", data,
            [
                ("reference", "HubSpot Deal ID"), ("name", "Deal"), ("customer", "Company"),
                ("stage", "Stage"), ("probability", "Probability"), ("status", "Status"),
                ("practice", "TECH"), ("region", "Region"), ("industry", "Industry"),
                ("owners", "Owner"),
            ],
        ),
    }  # fmt: skip


def to_xlsx(data: ExportData) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    summary.append([safe_text(data.title)])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    for key, value in _summary_rows(data):
        summary.append([key, safe_text(value)])
        summary.cell(summary.max_row, 1).font = Font(bold=True)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 90

    for table in tables(data).values():
        sheet = workbook.create_sheet(table.name)
        sheet.append(table.header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for row in table.rows:
            sheet.append(row)
        for column, header in enumerate(table.header, start=1):
            letter = get_column_letter(column)
            numeric = any(isinstance(r[column - 1], Decimal) for r in table.rows)
            sheet.column_dimensions[letter].width = (
                14 if numeric else max(12, min(40, len(header) + 4))
            )
            for cell in sheet[letter][1:]:
                if isinstance(cell.value, Decimal):
                    cell.number_format = "0.0000" if header == "Probability" else "#,##0.00"
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_csv(data: ExportData, view: str) -> bytes:
    """CSV: the export facts first (key, value), a blank line, then the chosen table. UTF-8 with
    a byte-order mark so Excel reads accented names correctly."""
    table = tables(data)[view]
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    for key, value in _summary_rows(data):
        writer.writerow([key, safe_text(value)])
    writer.writerow(["View", table.name])
    writer.writerow([])
    writer.writerow(table.header)
    for row in table.rows:
        writer.writerow(["" if v is None else v for v in row])
    return buffer.getvalue().encode("utf-8-sig")
