"""The HubSpot workbook adapter: explicit mappings, validation findings, nothing invented."""

# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false
# pyright: reportUnknownArgumentType=false, reportMissingTypeStubs=false

import hashlib
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from app.pipeline.source import (
    DEFAULT_WORKBOOK,
    HubSpotApiSource,
    SourceError,
    SourceSnapshot,
    WorkbookDealSource,
    build_source,
)

HEADERS = [
    "Record ID",
    "Deal Name",
    "Associated Company (Primary)",
    "Amount in company currency",
    "Deal owner",
    "Deal Stage",
    "Close Date",
    "Associated Company IDs (Primary)",
    "TECH (proposed)",
    "Region (proposed)",
    "Expected Start (proposed)",
    "Duration months (proposed)",
]


def deal_row(record_id: Any = "1001", **over: Any) -> list[Any]:
    values: dict[str, Any] = {
        "Record ID": record_id,
        "Deal Name": "Initech - Phase 2",
        "Associated Company (Primary)": "Initech",
        "Amount in company currency": 120000.0,
        "Deal owner": "Riley Shah",
        "Deal Stage": "Negotiation/Review",
        "Close Date": "2026-12-01 10:30",
        "Associated Company IDs (Primary)": "9000001",
        "TECH (proposed)": "APPIAN",
        "Region (proposed)": "NA",
        "Expected Start (proposed)": datetime(2026, 11, 1),
        "Duration months (proposed)": 6.0,
    }
    values.update(over)
    return [values[h] for h in HEADERS]


def build(
    path: Path,
    open_rows: list[list[Any]] | None = None,
    closed_rows: list[list[Any]] | None = None,
    *,
    stages: list[list[Any]] | None = None,
    links: list[list[Any]] | None = None,
    drop: tuple[str, ...] = (),
) -> Path:
    wb = openpyxl.Workbook()
    readme = wb.active
    assert readme is not None
    readme.title = "README"
    readme.append(["HubSpot test data"])
    readme.append(["As of", "2026-10-01. Test."])
    for name, rows in (("On_Deck", open_rows), ("Recent_Closed", closed_rows)):
        sheet = wb.create_sheet(name)
        sheet.append(HEADERS)
        for row in rows or []:
            sheet.append(row)
    stage_sheet = wb.create_sheet("Stage_Probabilities")
    stage_sheet.append(["Deal Stage", "Probability", "In forecast?", "Note"])
    for row in stages or [
        ["Presentation", 0.2, "Yes", "Placeholder"],
        ["Negotiation/Review", 0.5, "Yes", None],
        ["Closed Won", 1.0, "Committed", None],
        ["Closed Lost", 0.0, "No", None],
    ]:
        stage_sheet.append(row)
    link_sheet = wb.create_sheet("Deal_ChosenOne_Map")
    link_sheet.append(["Record ID", "Deal Name", "Stage", "Chosen One Row ID"])
    for row in links or []:
        link_sheet.append(row)
    for name in drop:
        del wb[name]
    wb.save(path)
    return path


def fetch(path: Path) -> SourceSnapshot:
    return WorkbookDealSource(path).fetch()


def codes(snapshot: SourceSnapshot, deal_id: str | None = None) -> list[str]:
    return [f.code for f in snapshot.findings if deal_id is None or f.deal_id == deal_id]


def test_a_clean_deal_is_mapped_field_by_field(tmp_path: Path) -> None:
    snap = fetch(build(tmp_path / "h.xlsx", [deal_row()]))

    (deal,) = snap.deals
    assert (deal.deal_id, deal.name, deal.stage, deal.company_name) == (
        "1001",
        "Initech - Phase 2",
        "Negotiation/Review",
        "Initech",
    )
    assert (deal.company_id, deal.owner_name, deal.practice, deal.region) == (
        "9000001",
        "Riley Shah",
        "APPIAN",
        "NA",
    )
    assert deal.amount == Decimal("120000.0") and deal.duration_months == 6
    assert deal.expected_start == date(2026, 11, 1)
    assert deal.close_date == datetime(2026, 12, 1, 10, 30, tzinfo=UTC)
    assert snap.findings == [] and snap.as_of == date(2026, 10, 1) and snap.complete
    assert deal.raw["Record ID"] == "1001"  # the original row is kept for traceability


def test_stage_probabilities_and_forecast_stage_rules_come_from_the_workbook(
    tmp_path: Path,
) -> None:
    snap = fetch(build(tmp_path / "h.xlsx", [deal_row()]))
    stages = {s.name: (s.probability, s.kind, s.in_forecast) for s in snap.stages}
    assert stages["Presentation"] == (Decimal("0.2"), "open", True)
    assert stages["Closed Won"] == (Decimal("1.0"), "won", True)
    assert stages["Closed Lost"] == (Decimal("0.0"), "lost", False)


def test_recent_closed_deals_are_read_alongside_open_ones(tmp_path: Path) -> None:
    snap = fetch(
        build(tmp_path / "h.xlsx", [deal_row("1")], [deal_row("2", **{"Deal Stage": "Closed Won"})])
    )
    assert {(d.deal_id, d.sheet) for d in snap.deals} == {("1", "On_Deck"), ("2", "Recent_Closed")}


def test_whitespace_is_trimmed_and_numeric_ids_become_text(tmp_path: Path) -> None:
    snap = fetch(
        build(
            tmp_path / "h.xlsx",
            [deal_row(60866525874, **{"Deal Name": "  Brambleton  -  Data Strategy "})],
        )
    )
    assert snap.deals[0].deal_id == "60866525874"
    assert snap.deals[0].name == "Brambleton - Data Strategy"


def test_company_id_in_the_name_column_is_resolved_by_id_not_by_name(tmp_path: Path) -> None:
    snap = fetch(
        build(
            tmp_path / "h.xlsx",
            [
                deal_row(
                    **{
                        "Associated Company (Primary)": "29146887085",
                        "Associated Company IDs (Primary)": "29146887085",
                    }
                )
            ],
        )
    )
    deal = snap.deals[0]
    assert deal.company_name is None and deal.company_id == "29146887085"
    assert codes(snap) == ["company_name_is_id"]


@pytest.mark.parametrize("missing", ["Expected Start (proposed)", "Duration months (proposed)"])
def test_missing_phasing_fields_are_reported_not_invented(tmp_path: Path, missing: str) -> None:
    snap = fetch(build(tmp_path / "h.xlsx", [deal_row(**{missing: None})]))
    deal = snap.deals[0]
    assert (deal.expected_start is None) or (deal.duration_months is None)
    assert codes(snap) == ["missing_start_or_duration"]


@pytest.mark.parametrize(
    ("bad", "code"),
    [
        ("N/A", "invalid_amount"),
        (-5.0, "invalid_amount"),
        (1.23456, "invalid_amount"),
        (None, "missing_amount"),
    ],
)
def test_unusable_amounts_are_findings_and_the_deal_has_no_amount(
    tmp_path: Path, bad: Any, code: str
) -> None:
    snap = fetch(build(tmp_path / "h.xlsx", [deal_row(**{"Amount in company currency": bad})]))
    assert snap.deals[0].amount is None
    assert code in codes(snap)


def test_four_decimal_amounts_are_kept_exactly(tmp_path: Path) -> None:
    snap = fetch(
        build(tmp_path / "h.xlsx", [deal_row(**{"Amount in company currency": 308409.0508})])
    )
    assert snap.deals[0].amount == Decimal("308409.0508")


@pytest.mark.parametrize("bad", [4.5, 0, -2, "six"])
def test_durations_must_be_whole_positive_months(tmp_path: Path, bad: Any) -> None:
    snap = fetch(build(tmp_path / "h.xlsx", [deal_row(**{"Duration months (proposed)": bad})]))
    assert snap.deals[0].duration_months is None
    assert "invalid_duration" in codes(snap)


def test_close_dates_that_cannot_be_read_are_reported(tmp_path: Path) -> None:
    snap = fetch(build(tmp_path / "h.xlsx", [deal_row(**{"Close Date": "next Tuesday"})]))
    assert snap.deals[0].close_date is None and "invalid_close_date" in codes(snap)


def test_duplicate_and_blank_record_ids_are_rejected_with_findings(tmp_path: Path) -> None:
    snap = fetch(
        build(
            tmp_path / "h.xlsx",
            [deal_row("1"), deal_row("1", **{"Deal Name": "Again"}), deal_row(None)],
        )
    )
    assert [d.name for d in snap.deals] == ["Initech - Phase 2"]
    assert sorted(codes(snap)) == ["duplicate_deal_id", "missing_deal_id"]


def test_rows_without_a_name_or_stage_are_skipped(tmp_path: Path) -> None:
    snap = fetch(build(tmp_path / "h.xlsx", [deal_row("1", **{"Deal Stage": None})]))
    assert snap.deals == [] and codes(snap) == ["missing_required_field"]


def test_explicit_links_are_read_and_nothing_is_inferred(tmp_path: Path) -> None:
    snap = fetch(
        build(tmp_path / "h.xlsx", [deal_row("1"), deal_row("2")], links=[["1", "x", "y", "R1007"]])
    )
    assert [(link.deal_id, link.project_row_id) for link in snap.links] == [("1", "R1007")]


@pytest.mark.parametrize("sheet", ["On_Deck", "Stage_Probabilities"])
def test_a_missing_required_sheet_is_a_source_error(tmp_path: Path, sheet: str) -> None:
    with pytest.raises(SourceError, match=sheet):
        fetch(build(tmp_path / "h.xlsx", [deal_row()], drop=(sheet,)))


def test_a_missing_required_column_is_a_source_error(tmp_path: Path) -> None:
    path = build(tmp_path / "h.xlsx", [deal_row()])
    wb = openpyxl.load_workbook(path)
    wb["On_Deck"]["F1"].value = "Stage"  # renames 'Deal Stage'
    wb.save(path)
    with pytest.raises(SourceError, match="Deal Stage"):
        fetch(path)


def test_a_stage_probability_outside_zero_to_one_is_a_source_error(tmp_path: Path) -> None:
    path = build(tmp_path / "h.xlsx", [deal_row()], stages=[["Presentation", 1.5, "Yes", None]])
    with pytest.raises(SourceError, match="probability"):
        fetch(path)


def test_unreadable_files_are_a_source_error(tmp_path: Path) -> None:
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"this is not a workbook")
    with pytest.raises(SourceError, match="Cannot read"):
        fetch(broken)
    with pytest.raises(SourceError, match="Cannot read"):
        fetch(tmp_path / "missing.xlsx")


def test_reading_never_modifies_the_workbook(tmp_path: Path) -> None:
    path = build(tmp_path / "h.xlsx", [deal_row()])
    before = (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
    fetch(path)
    assert (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns) == before


def test_the_live_adapter_is_unavailable_until_implemented_and_never_writes() -> None:
    with pytest.raises(SourceError, match="HUBSPOT_ACCESS_TOKEN"):
        HubSpotApiSource(None).fetch()
    with pytest.raises(SourceError, match="not implemented"):
        HubSpotApiSource("token").fetch()
    # The adapter contract is read-only: a source can fetch, and nothing else.
    assert [n for n in dir(HubSpotApiSource) if not n.startswith("_")] == ["fetch", "name"]


def test_build_source_follows_the_configuration(tmp_path: Path) -> None:
    assert isinstance(build_source("workbook", str(tmp_path / "x.xlsx"), None), WorkbookDealSource)
    assert isinstance(build_source("api", None, "t"), HubSpotApiSource)
    default = build_source("workbook", None, None)
    assert isinstance(default, WorkbookDealSource) and default.path == DEFAULT_WORKBOOK


@pytest.mark.skipif(not DEFAULT_WORKBOOK.is_file(), reason="reference workbook not present")
def test_the_supplied_hubspot_export_is_read_completely() -> None:
    snap = WorkbookDealSource(DEFAULT_WORKBOOK).fetch()
    assert len(snap.deals) == 53 and len(snap.stages) == 6 and snap.as_of == date(2026, 10, 1)
    assert [(link.deal_id, link.project_row_id) for link in snap.links] == [
        ("57018801551", "R1007"),
        ("64360839985", "R1003"),
    ]
    assert sorted(codes(snap)) == ["company_name_is_id"] * 3 + ["missing_start_or_duration"] * 2
    assert {d.company_id for d in snap.deals if d.company_name is None} == {"29146887085"}
