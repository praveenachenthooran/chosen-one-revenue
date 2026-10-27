"""Read and validate the Chosen One Excel workbook. Pure parsing: no database access.

Columns are located by header name, never by fixed position. Source financial amounts
are never rounded, coerced or repaired: a cell is either accepted exactly as written or
reported, and a row with any invalid amount is rejected as a whole.
"""

# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false
# pyright: reportUnknownArgumentType=false, reportMissingTypeStubs=false

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

FORECAST_SHEET = "Forecast"
USERS_SHEET = "Users"
README_SHEET = "README"
SUMMARY_SHEET = "Summary"

VALID_CATEGORIES = ("EE", "EN", "NN")
CATEGORY_DESCRIPTIONS = {
    # PRD section 9 marks these as "[confirm]".
    "EE": "Existing customer / existing work",
    "EN": "Existing customer / new work",
    "NN": "New customer / new work",
}
NOT_APPLICABLE = {"N/A", "#N/A"}  # "NA" is a real value: Region = North America
NEAR_DUPLICATE_RATIO = 0.92

REQUIRED_HEADERS = ("Customer", "Project", "Row ID", "Status")
OPTIONAL_TEXT_HEADERS = {
    "practice": "TECH",
    "region": "Region",
    "industry": "Industry",
    "country": "Country",
    "billing_type": "Billing Type",
}
OWNER_HEADERS = {
    "client_partner": "Client Partner",
    "ae": "AE",
    "commercial_owner": "Commercial Owner",
}


class WorkbookStructureError(Exception):
    """The workbook does not have the expected sheets or headers; nothing can be imported."""


@dataclass(frozen=True)
class SourceIssue:
    severity: str  # "error" (row rejected) or "warning" (row imported)
    code: str
    sheet: str
    message: str
    row_number: int | None = None
    source_row_id: str | None = None
    field: str | None = None
    value: str | None = None
    row_rejected: bool = False
    details: dict[str, Any] | None = None


@dataclass
class ParsedUser:
    code: str
    name: str
    roles: list[str]
    email: str | None
    location: str | None
    notes: str | None


@dataclass
class ParsedProject:
    row_number: int
    source_row_id: str
    customer: str
    name: str
    status: str
    owner_codes: dict[str, str]  # role -> raw workbook code
    practice: str | None
    region: str | None
    industry: str | None
    country: str | None
    billing_type: str | None
    category: str | None
    last_updated: date | None
    last_updated_by: str | None
    months: dict[date, Decimal] = field(default_factory=lambda: {})
    cells: dict[date, str] = field(default_factory=lambda: {})


@dataclass
class ParsedWorkbook:
    file_name: str
    sha256: str
    as_of: date | None
    month_columns: list[date]
    projects: list[ParsedProject]
    users: list[ParsedUser]
    issues: list[SourceIssue]
    rows_read: int
    blank_rows: int
    # Independent per-month sum of every numeric cell on every project row (rejected rows
    # included), so a rejected row shows up as a variance instead of disappearing.
    source_totals: dict[date, Decimal]
    # Totals Excel itself last calculated (cached SUBTOTAL / Summary values), if any.
    excel_totals: dict[date, Decimal | None]

    @property
    def rejected(self) -> list[SourceIssue]:
        return [i for i in self.issues if i.row_rejected]

    @property
    def rejected_row_count(self) -> int:
        return len({i.row_number for i in self.issues if i.row_rejected})


class _RowReport:
    """Collects the errors and warnings found on one source row."""

    def __init__(self, row: int, source_id: str | None) -> None:
        self.row = row
        self.source_id = source_id
        self.errors: list[SourceIssue] = []
        self.warnings: list[SourceIssue] = []

    def _issue(self, severity: str, code: str, message: str, header: str | None,
               value: object, rejected: bool) -> SourceIssue:  # fmt: skip
        return SourceIssue(severity, code, FORECAST_SHEET, message, self.row, self.source_id,
                           header, None if value is None else str(value),
                           row_rejected=rejected)  # fmt: skip

    def reject(
        self, code: str, message: str, header: str | None = None, value: object = None
    ) -> None:
        self.errors.append(self._issue("error", code, message, header, value, True))

    def warn(
        self, code: str, message: str, header: str | None = None, value: object = None
    ) -> None:
        self.warnings.append(self._issue("warning", code, message, header, value, False))


def _clean(value: object) -> str | None:
    """Trim a text cell; blank cells become None. Non-text values are stringified."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _text_or_none(value: object) -> str | None:
    """Like _clean, but the workbook's 'N/A' placeholders also become None."""
    text = _clean(value)
    if text is None or text.upper() in NOT_APPLICABLE:
        return None
    return text


def _normalise_name(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _to_month(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _parse_amount(value: object) -> tuple[Decimal | None, str | None]:
    """Return (amount, error_code). A blank cell returns (None, None)."""
    if value is None:
        return None, None
    if isinstance(value, bool):
        return None, "invalid_amount"
    if isinstance(value, int):
        amount = Decimal(value)
    elif isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None, "invalid_amount"
        amount = Decimal(repr(value))  # shortest round-trip text, no binary-float noise
    elif isinstance(value, Decimal):
        amount = value
    else:
        return None, "invalid_amount"  # text such as "N/A" or "1,200"
    if amount < 0:
        return None, "negative_amount"
    if amount != amount.quantize(Decimal("0.01")):
        return None, "amount_precision"  # would need rounding; refuse rather than alter
    return amount.quantize(Decimal("0.01")), None


def _find_header_row(sheet: Worksheet) -> int:
    for row in range(1, min(sheet.max_row, 15) + 1):
        values = {_clean(sheet.cell(row, col).value) for col in range(1, sheet.max_column + 1)}
        if {"Customer", "Project"} <= values:
            return row
    raise WorkbookStructureError(
        f"Sheet '{sheet.title}': no header row containing 'Customer' and 'Project' found"
    )


def _read_as_of(workbook: openpyxl.Workbook) -> date | None:
    if README_SHEET not in workbook.sheetnames:
        return None
    sheet = workbook[README_SHEET]
    for row in sheet.iter_rows(values_only=True):
        if row and _clean(row[0]) == "As of" and len(row) > 1 and row[1]:
            match = re.search(r"\d{4}-\d{2}-\d{2}", str(row[1]))
            if match:
                return date.fromisoformat(match.group())
    return None


def _parse_users(workbook: openpyxl.Workbook, issues: list[SourceIssue]) -> list[ParsedUser]:
    if USERS_SHEET not in workbook.sheetnames:
        issues.append(
            SourceIssue(
                "warning", "users_sheet_missing", USERS_SHEET, "No Users sheet; owners unresolved"
            )
        )
        return []
    sheet = workbook[USERS_SHEET]
    headers = {_clean(c.value): c.column for c in sheet[1] if _clean(c.value)}
    missing = [h for h in ("Code", "Name") if h not in headers]
    if missing:
        raise WorkbookStructureError(f"Users sheet is missing header(s): {', '.join(missing)}")

    def cell(row: int, header: str) -> object:
        column = headers.get(header)
        return sheet.cell(row, column).value if column else None

    users: list[ParsedUser] = []
    seen_codes: set[str] = set()
    seen_emails: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        code, name = _clean(cell(row, "Code")), _clean(cell(row, "Name"))
        if code is None and name is None:
            continue
        if code is None or name is None:
            issues.append(
                SourceIssue(
                    "error", "user_missing_required_field", USERS_SHEET,
                    "User row needs both Code and Name; skipped",
                    row_number=row, row_rejected=True,
                )
            )  # fmt: skip
            continue
        if code in seen_codes:
            issues.append(
                SourceIssue(
                    "error", "duplicate_user_code", USERS_SHEET,
                    f"Duplicate user code '{code}'; later row skipped",
                    row_number=row, field="Code", value=code, row_rejected=True,
                )
            )  # fmt: skip
            continue
        seen_codes.add(code)
        email = _clean(cell(row, "Email"))
        if email is not None:
            email = email.lower()
            if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email) or email in seen_emails:
                issues.append(
                    SourceIssue(
                        "warning", "user_email_invalid_or_duplicate", USERS_SHEET,
                        f"Email for '{code}' is malformed or duplicated; imported without email",
                        row_number=row, field="Email", value=email,
                    )
                )  # fmt: skip
                email = None
            else:
                seen_emails.add(email)
        roles_text = _clean(cell(row, "Role(s)")) or ""
        roles = [r.strip() for r in roles_text.split(";") if r.strip()]
        users.append(
            ParsedUser(code, name, roles, email, _clean(cell(row, "Location")),
                       _clean(cell(row, "Notes")))
        )  # fmt: skip
    return users


def _month_columns(sheet: Worksheet, header_row: int) -> dict[int, date]:
    columns: dict[int, date] = {}
    seen: set[date] = set()
    for col in range(1, sheet.max_column + 1):
        month = _to_month(sheet.cell(header_row, col).value)
        if month is None:
            continue
        if month.day != 1:
            raise WorkbookStructureError(
                f"Forecast month header in column {col} is {month}, not the first of a month"
            )
        if month in seen:
            raise WorkbookStructureError(f"Forecast month header {month:%Y-%m} appears twice")
        seen.add(month)
        columns[col] = month
    if not columns:
        raise WorkbookStructureError("Forecast sheet has no date-typed month column headers")
    return columns


def _cached_totals(
    formulas: openpyxl.Workbook, values: openpyxl.Workbook, header_row: int, months: dict[int, date]
) -> dict[date, Decimal | None]:
    """Totals Excel last calculated: Forecast SUBTOTAL row, else the Summary 'Total' row."""
    totals: dict[date, Decimal | None] = dict.fromkeys(months.values())

    def as_decimal(value: object) -> Decimal | None:
        if isinstance(value, bool) or not isinstance(value, int | float | Decimal):
            return None
        return Decimal(repr(value)) if isinstance(value, float) else Decimal(value)

    forecast_values = values[FORECAST_SHEET]
    for row in range(1, header_row):
        for col, month in months.items():
            formula = formulas[FORECAST_SHEET].cell(row, col).value
            if isinstance(formula, str) and formula.upper().startswith("=SUBTOTAL(9,"):
                totals[month] = as_decimal(forecast_values.cell(row, col).value)
    if SUMMARY_SHEET in values.sheetnames and any(v is None for v in totals.values()):
        summary = values[SUMMARY_SHEET]
        header = next(
            (
                r
                for r in range(1, summary.max_row + 1)
                if _clean(summary.cell(r, 1).value) == "Group"
            ),
            None,
        )
        total_row = next(
            (
                r
                for r in range(1, summary.max_row + 1)
                if _clean(summary.cell(r, 1).value) == "Total"
            ),
            None,
        )
        if header and total_row:
            for col in range(2, summary.max_column + 1):
                month = _to_month(summary.cell(header, col).value)
                if month is not None and month in totals and totals[month] is None:
                    totals[month] = as_decimal(summary.cell(total_row, col).value)
    return totals


def parse_workbook(path: Path) -> ParsedWorkbook:
    """Parse the workbook at `path` without modifying it."""
    raw = path.read_bytes()
    sha256 = hashlib.sha256(raw).hexdigest()
    formulas = openpyxl.load_workbook(BytesIO(raw), data_only=False)
    values = openpyxl.load_workbook(BytesIO(raw), data_only=True)
    if FORECAST_SHEET not in formulas.sheetnames:
        raise WorkbookStructureError(f"Workbook has no '{FORECAST_SHEET}' sheet")

    issues: list[SourceIssue] = []
    users = _parse_users(formulas, issues)
    known_codes = {u.code for u in users}

    sheet = formulas[FORECAST_SHEET]
    header_row = _find_header_row(sheet)
    headers = {_clean(sheet.cell(header_row, c).value): c for c in range(1, sheet.max_column + 1)}
    headers.pop(None, None)
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise WorkbookStructureError(f"Forecast sheet is missing header(s): {', '.join(missing)}")
    months = _month_columns(sheet, header_row)

    projects: list[ParsedProject] = []
    source_totals: dict[date, Decimal] = dict.fromkeys(months.values(), Decimal("0.00"))
    seen_ids: dict[str, int] = {}
    seen_names: dict[tuple[str, str], tuple[int, str]] = {}
    by_customer: dict[str, list[tuple[str, str, int]]] = {}
    rows_read = blank_rows = 0

    def text_cell(row: int, header: str) -> object:
        col = headers.get(header)
        return sheet.cell(row, col).value if col else None

    for row in range(header_row + 1, sheet.max_row + 1):
        cells = [sheet.cell(row, c) for c in range(1, sheet.max_column + 1)]
        if all(c.value is None or _clean(c.value) is None for c in cells):
            blank_rows += 1
            continue
        rows_read += 1
        customer, name = _clean(text_cell(row, "Customer")), _clean(text_cell(row, "Project"))
        source_id, status = _clean(text_cell(row, "Row ID")), _clean(text_cell(row, "Status"))
        report = _RowReport(row, source_id)
        reject, warn = report.reject, report.warn

        for header, value in (("Customer", customer), ("Project", name),
                              ("Row ID", source_id), ("Status", status)):  # fmt: skip
            if value is None:
                reject("missing_required_field", f"Required field '{header}' is blank", header)

        # Monthly amounts: every cell is validated; any failure rejects the whole row.
        amounts: dict[date, Decimal] = {}
        cell_refs: dict[date, str] = {}
        for col, month in months.items():
            cell = sheet.cell(row, col)
            if cell.data_type == "f":
                reject("formula_in_amount",
                       f"{cell.coordinate} contains a formula; amounts must be entered values",
                       f"{month:%Y-%m}", cell.value)  # fmt: skip
                continue
            amount, error = _parse_amount(cell.value)
            if error:
                messages = {
                    "invalid_amount": f"{cell.coordinate} is not a number",
                    "negative_amount": f"{cell.coordinate} is negative",
                    "amount_precision": f"{cell.coordinate} has more than 2 decimal places",
                }
                reject(error, messages[error], f"{month:%Y-%m}", cell.value)
            if isinstance(cell.value, int | float | Decimal) and not isinstance(cell.value, bool):
                source_totals[month] += Decimal(repr(cell.value)) if isinstance(
                    cell.value, float) else Decimal(cell.value)  # fmt: skip
            if amount is not None:
                amounts[month] = amount
                cell_refs[month] = f"{FORECAST_SHEET}!{cell.coordinate}"

        category = _clean(text_cell(row, "CATG"))
        if category is not None and category.upper() not in VALID_CATEGORIES:
            reject("invalid_category", f"Category '{category}' is not one of EE, EN, NN",
                   "CATG", category)  # fmt: skip
        category = category.upper() if category else None

        if source_id is not None:
            if source_id in seen_ids:
                reject("duplicate_row_id",
                       f"Row ID '{source_id}' already used on row {seen_ids[source_id]}",
                       "Row ID", source_id)  # fmt: skip
            else:
                seen_ids[source_id] = row
        if customer is not None and name is not None:
            key = (_normalise_name(customer), _normalise_name(name))
            if key in seen_names:
                first_row, first_id = seen_names[key]
                reject("duplicate_project",
                       f"Same customer and project as row {first_row} ({first_id})",
                       "Project", name)  # fmt: skip
            else:
                seen_names[key] = (row, source_id or "")

        if report.errors:
            issues.extend(report.errors)
            continue
        assert customer and name and source_id and status

        owner_codes: dict[str, str] = {}
        for role, header in OWNER_HEADERS.items():
            raw_code = _text_or_none(text_cell(row, header))
            if raw_code is None:
                if role == "client_partner":
                    warn("missing_client_partner",
                         "No Client Partner (blank or N/A); every project needs an owner",
                         header, _clean(text_cell(row, header)))  # fmt: skip
            elif raw_code not in known_codes:
                warn("unknown_owner_code",
                     f"{header} '{raw_code}' is not a code on the Users sheet; owner not linked",
                     header, raw_code)  # fmt: skip
            else:
                owner_codes[role] = raw_code
        optional: dict[str, str | None] = {}
        for attr, header in OPTIONAL_TEXT_HEADERS.items():
            optional[attr] = _text_or_none(text_cell(row, header))
            if optional[attr] is None and attr != "country" and attr != "billing_type":
                warn(f"missing_{attr}", f"'{header}' is blank", header)
        if category is None:
            warn("missing_category", "Category (CATG) is blank", "CATG")
        if status not in ("Active", "Completed"):
            warn("unrecognized_status", f"Status '{status}' is not Active or Completed",
                 "Status", status)  # fmt: skip
        if not amounts:
            warn("no_revenue", "Project has no monthly revenue in any fiscal month")

        raw_updated = text_cell(row, "Last Updated")
        last_updated = _to_month(raw_updated)
        if raw_updated is not None and last_updated is None:
            warn("invalid_last_updated", "'Last Updated' is not a date; ignored",
                 "Last Updated", raw_updated)  # fmt: skip

        for other_name, other_id, other_row in by_customer.get(_normalise_name(customer), []):
            ratio = SequenceMatcher(
                None, _normalise_name(name), _normalise_name(other_name)
            ).ratio()
            if ratio >= NEAR_DUPLICATE_RATIO:
                report.warnings.append(
                    SourceIssue(
                        "warning", "possible_duplicate_project", FORECAST_SHEET,
                        f"Project name is very close to '{other_name}' ({other_id}, "
                        f"row {other_row}) for the same customer; check for a double entry",
                        row, source_id, "Project", name,
                        details={"similar_to_row_id": other_id, "similarity": round(ratio, 3)},
                    )
                )  # fmt: skip
        by_customer.setdefault(_normalise_name(customer), []).append((name, source_id, row))

        issues.extend(report.warnings)
        projects.append(
            ParsedProject(
                row_number=row, source_row_id=source_id, customer=customer, name=name,
                status=status, owner_codes=owner_codes,
                practice=optional["practice"], region=optional["region"],
                industry=optional["industry"], country=optional["country"],
                billing_type=optional["billing_type"], category=category,
                last_updated=last_updated,
                last_updated_by=_clean(text_cell(row, "Last Updated By")),
                months=amounts, cells=cell_refs,
            )
        )  # fmt: skip

    return ParsedWorkbook(
        file_name=path.name,
        sha256=sha256,
        as_of=_read_as_of(formulas),
        month_columns=sorted(months.values()),
        projects=projects,
        users=users,
        issues=issues,
        rows_read=rows_read,
        blank_rows=blank_rows,
        source_totals=source_totals,
        excel_totals=_cached_totals(formulas, values, header_row, months),
    )
